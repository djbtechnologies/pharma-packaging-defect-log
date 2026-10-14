"""
Rebuild the analysis sheets and charts in packaging_defect_log.xlsx.

Each recall's final category is its Reviewed Category if one was entered,
otherwise its Defect Category. Every sheet, chart, and printed number uses
that final category.

Sheets: Pareto, Trend, By Class, Injectables, and Review Check (added once
any Reviewed Category is filled in). The Pareto and Trend charts are also
saved as PNGs in charts/.

Close the workbook in Excel before running, since Excel locks the file.

Usage: python analyze.py
"""

import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.ticker import PercentFormatter
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from categories import CATEGORY_NAMES, EXCLUDE, WORKBOOK, add_dropdowns

OUTPUT_SHEETS = ["Pareto", "Trend", "By Class", "Injectables", "Review Check"]
CHART_DIR = Path("charts")
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BLUE, ORANGE, AQUA = "2a78d6", "eb6834", "1baf7a"  # colorblind-safe trio
INK, MUTED, GRID = "#0b0b0b", "#52514e", "#e4e3df"


def read_log(wb):
    rows = list(wb["Recall Log"].iter_rows(values_only=True))
    df = pd.DataFrame(rows[1:], columns=rows[0])
    df = df[df["Recall Number"].notna()]
    df["Reviewed Category"] = df["Reviewed Category"].fillna("").astype(str).str.strip()

    bad = df[(df["Reviewed Category"] != "") & ~df["Reviewed Category"].isin(CATEGORY_NAMES)]
    if len(bad):
        print("These Reviewed Category values aren't on the Categories sheet. Fix them and rerun:")
        for _, r in bad.iterrows():
            print(f"  {r['Recall Number']}: '{r['Reviewed Category']}'")
        raise SystemExit(1)

    reviewed = df["Reviewed Category"] != ""
    df["Final Category"] = df["Reviewed Category"].where(reviewed, df["Defect Category"])
    df["Year"] = df["Year"].astype(int)
    return df


def pareto_table(df):
    counts = df.loc[df["Final Category"] != EXCLUDE, "Final Category"].value_counts()
    t = pd.DataFrame({"Category": counts.index, "Count": counts.values})
    t["Percent"] = t["Count"] / t["Count"].sum()
    t["Cumulative Percent"] = t["Percent"].cumsum()
    return t


# ---------- sheet helpers ----------

def write_table(ws, start_row, headers, rows, percent_cols=()):
    """Write a bold header plus rows starting at start_row. Returns the last row written."""
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    r = start_row
    for values in rows:
        r += 1
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if headers[c - 1] in percent_cols:
                cell.number_format = "0.0%"
    return r


def note(ws, row, text, title=False):
    ws.cell(row=row, column=1, value=text).font = Font(bold=title, size=13 if title else 11)


def widths(ws, w):
    for col, width in w.items():
        ws.column_dimensions[col].width = width


def style_series(series, color):
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 28000  # about 2pt
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color
    series.smooth = False


# ---------- analysis sheets ----------

def pareto_sheet(wb, df, name, subtitle):
    ws = wb.create_sheet(name)
    t = pareto_table(df)
    excluded = int((df["Final Category"] == EXCLUDE).sum())

    note(ws, 1, f"{name}: packaging recalls by defect category", title=True)
    note(ws, 2, subtitle)
    note(ws, 3, f"{int(t['Count'].sum())} packaging recalls. {excluded} marked '{EXCLUDE}' are not counted.")
    top = 5
    last = write_table(ws, top, list(t.columns), t.values.tolist(),
                       percent_cols=("Percent", "Cumulative Percent"))
    widths(ws, {"A": 36, "B": 9, "C": 10, "D": 13})

    cats = Reference(ws, min_col=1, min_row=top + 1, max_row=last)

    bars = BarChart()
    bars.type = "col"
    bars.title = f"{name} of packaging recall categories"
    bars.y_axis.title = "Recalls"
    bars.add_data(Reference(ws, min_col=2, min_row=top, max_row=last), titles_from_data=True)
    bars.set_categories(cats)
    bars.series[0].graphicalProperties.solidFill = BLUE
    bars.series[0].graphicalProperties.line.solidFill = BLUE
    bars.gapWidth = 60

    # Cumulative percent on a secondary 0-100% axis on the right
    line = LineChart()
    line.add_data(Reference(ws, min_col=4, min_row=top, max_row=last), titles_from_data=True)
    line.set_categories(cats)
    line.y_axis.axId = 200
    line.y_axis.title = "Cumulative percent"
    line.y_axis.number_format = "0%"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 1
    line.y_axis.crosses = "max"
    line.y_axis.majorGridlines = None
    style_series(line.series[0], ORANGE)

    # openpyxl 3.1 hides chart axes unless delete is set to False
    bars.x_axis.delete = False
    bars.y_axis.delete = False
    line.y_axis.delete = False

    bars += line
    bars.legend.position = "b"
    bars.width, bars.height = 26, 13
    ws.add_chart(bars, "F2")
    return t


def trend_sheet(wb, df, top3):
    ws = wb.create_sheet("Trend")
    years = list(range(df["Year"].min(), df["Year"].max() + 1))
    counts = (df[df["Final Category"].isin(top3)]
              .groupby(["Year", "Final Category"]).size().unstack(fill_value=0)
              .reindex(index=years, columns=top3, fill_value=0))

    note(ws, 1, "Trend: recalls per year for the top 3 packaging categories", title=True)
    note(ws, 2, f"Top 3 from the Pareto sheet. {years[-1]} is a partial year.")
    top = 4
    last = write_table(ws, top, ["Year"] + top3, [[y] + counts.loc[y].tolist() for y in years])
    widths(ws, {"A": 8, "B": 22, "C": 22, "D": 22})

    chart = LineChart()
    chart.title = "Top 3 packaging failure modes per year"
    chart.y_axis.title = "Recalls"
    chart.x_axis.title = "Report year"
    chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=top, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    for series, color in zip(chart.series, [BLUE, ORANGE, AQUA]):
        style_series(series, color)
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    chart.width, chart.height = 24, 12
    ws.add_chart(chart, "F2")
    return counts


def class_sheet(wb, df):
    ws = wb.create_sheet("By Class")
    pkg = df[df["Final Category"] != EXCLUDE]
    classes = ["Class I", "Class II", "Class III"]
    order = pareto_table(df)["Category"].tolist()
    t = pd.crosstab(pkg["Final Category"], pkg["Class"]).reindex(index=order, columns=classes, fill_value=0)

    note(ws, 1, "By Class: packaging recalls by category and FDA recall class", title=True)
    note(ws, 2, "Class I: could cause serious harm or death. Class II: temporary or reversible harm. "
                "Class III: unlikely to cause harm.")
    rows = [[c] + t.loc[c].tolist() + [int(t.loc[c].sum())] for c in order]
    rows.append(["Total"] + t.sum().astype(int).tolist() + [int(t.values.sum())])
    last = write_table(ws, 4, ["Category"] + classes + ["Total"], rows)
    for cell in ws[last]:
        cell.font = Font(bold=True)
    widths(ws, {"A": 36, "B": 10, "C": 10, "D": 10, "E": 10})


def review_sheet(wb, df):
    """How often the reviewed category matches the original Defect Category."""
    ws = wb.create_sheet("Review Check")
    rev = df[df["Reviewed Category"] != ""]
    match = rev["Reviewed Category"] == rev["Defect Category"]
    rate = match.mean()

    note(ws, 1, "Review Check: reviewed categories vs. original categories", title=True)
    write_table(ws, 3, ["Measure", "Value"], [
        ["Rows reviewed", len(rev)],
        ["Rows in the log", len(df)],
        ["Rows where the review matched", int(match.sum())],
        ["Agreement rate", rate],
    ])
    ws["B7"].number_format = "0.0%"

    diff = rev[~match]
    note(ws, 9, f"Disagreements ({len(diff)} rows)", title=True)
    write_table(ws, 10, ["Recall Number", "Original Category", "Reviewed Category", "Reason for Recall"],
                diff[["Recall Number", "Defect Category", "Reviewed Category", "Reason for Recall"]].values.tolist())
    widths(ws, {"A": 34, "B": 34, "C": 34, "D": 80})
    ws.freeze_panes = "A11"
    return len(rev), rate


# ---------- PNG charts ----------

def clean_axes(ax):
    for side in ["top", "right"]:
        ax.spines[side].set_visible(False)
    for side in ["left", "bottom"]:
        ax.spines[side].set_color(MUTED)
    ax.tick_params(colors=MUTED)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def pareto_png(t, path):
    fig, ax = plt.subplots(figsize=(11, 6))
    x = list(range(len(t)))
    bars = ax.bar(x, t["Count"], color=f"#{BLUE}", width=0.7)
    clean_axes(ax)
    ax.set_xticks(x, ["\n".join(textwrap.wrap(c, 13)) for c in t["Category"]], fontsize=9, color=INK)
    ax.set_ylabel("Recalls", color=MUTED)

    ax2 = ax.twinx()
    ax2.plot(x, t["Cumulative Percent"], color=f"#{ORANGE}", marker="o", linewidth=2, markersize=7)
    ax2.set_ylim(0, 1.05)
    ax2.yaxis.set_major_formatter(PercentFormatter(1.0))
    ax2.spines["top"].set_visible(False)
    ax2.spines["right"].set_color(MUTED)
    ax2.tick_params(colors=MUTED)
    ax2.set_ylabel("Cumulative percent", color=MUTED)
    ax2.axhline(0.8, color=MUTED, linestyle=":", linewidth=1)
    ax2.text(len(t) - 0.5, 0.81, "80%", color=MUTED, fontsize=9, ha="right")

    # Count labels go on ax2 (the top layer) with a white backing so the 80% line doesn't cross them
    for bar, n in zip(bars, t["Count"]):
        ax2.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 1, str(n),
                 transform=ax.transData, ha="center", va="bottom", fontsize=9, color=INK,
                 zorder=5, bbox=dict(facecolor="white", edgecolor="none", pad=1))

    ax.set_title("Pareto of packaging recall categories", loc="left", fontsize=13, color=INK)
    fig.legend([bars, ax2.lines[0]], ["Recalls (left axis)", "Cumulative percent (right axis)"],
               loc="upper center", bbox_to_anchor=(0.5, 0.02), ncol=2, frameon=False)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def trend_png(counts, path):
    fig, ax = plt.subplots(figsize=(11, 5.5))
    years = counts.index.tolist()
    for cat, color in zip(counts.columns, [BLUE, ORANGE, AQUA]):
        ax.plot(years, counts[cat], color=f"#{color}", marker="o", linewidth=2, markersize=7, label=cat)
    clean_axes(ax)
    ax.set_xticks(years)
    ax.set_ylabel("Recalls", color=MUTED)
    ax.set_xlabel(f"Report year ({years[-1]} is a partial year)", color=MUTED)
    ax.set_title("Top 3 packaging failure modes per year", loc="left", fontsize=13, color=INK)
    ax.legend(frameon=False, loc="upper right")
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def main():
    wb = load_workbook(WORKBOOK)
    df = read_log(wb)
    for name in OUTPUT_SHEETS:
        if name in wb.sheetnames:
            del wb[name]

    rule = "Uses Reviewed Category where filled in, otherwise Defect Category."
    pareto = pareto_sheet(wb, df, "Pareto", rule)
    top3 = pareto["Category"].head(3).tolist()
    trend = trend_sheet(wb, df, top3)
    class_sheet(wb, df)
    inj = pareto_sheet(wb, df[df["Injectable"] == "Yes"], "Injectables", rule + " Injectables only.")
    review = review_sheet(wb, df) if (df["Reviewed Category"] != "").any() else None

    add_dropdowns(wb)
    wb.save(WORKBOOK)

    CHART_DIR.mkdir(exist_ok=True)
    pareto_png(pareto, CHART_DIR / "pareto.png")
    trend_png(trend, CHART_DIR / "trend.png")

    print(f"{len(df)} recalls in the log, {int(pareto['Count'].sum())} packaging "
          f"({int(inj['Count'].sum())} injectable)\n")
    for _, r in pareto.iterrows():
        print(f"  {r['Category']:<38}{r['Count']:>4}  {r['Percent']:6.1%}   cum {r['Cumulative Percent']:6.1%}")
    print(f"\nTop 3 = {pareto['Cumulative Percent'].iloc[2]:.1%} of packaging recalls")
    if review:
        print(f"Review check: {review[0]} rows reviewed, {review[1]:.1%} agreement")


if __name__ == "__main__":
    main()
