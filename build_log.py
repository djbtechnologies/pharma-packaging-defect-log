"""
Build the Recall Log and Categories sheets of packaging_defect_log.xlsx.

Joins data/packaging_recalls.csv with data/defect_categories.csv. If the
workbook already exists, anything typed in Reviewed Category is carried over,
so rebuilding is safe.

Run analyze.py afterward to build the analysis sheets.

Usage: python build_log.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from categories import CATEGORIES, LOG_COLUMNS, WORKBOOK, add_dropdowns

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
WIDTHS = {
    "Recall Number": 14, "Event ID": 9, "Report Date": 12, "Year": 7, "Firm": 26,
    "Product": 45, "Route": 16, "Injectable": 10, "Class": 9, "Status": 11,
    "Reason for Recall": 60, "Defect Category": 30, "Category Rationale": 50,
    "Reviewed Category": 30,
}
WRAPPED = {"Firm", "Product", "Route", "Reason for Recall", "Defect Category",
           "Category Rationale", "Reviewed Category"}


def load_log():
    recalls = pd.read_csv("data/packaging_recalls.csv", dtype={"event_id": str})
    cats = pd.read_csv("data/defect_categories.csv")
    df = recalls.merge(cats, on="recall_number", how="left")
    missing = df.loc[df["defect_category"].isna(), "recall_number"].tolist()
    if missing:
        raise ValueError(f"no category for: {missing}")

    df["report_date"] = pd.to_datetime(df["report_date"])
    df = df.sort_values(["report_date", "recall_number"], ascending=[False, True])
    return pd.DataFrame({
        "Recall Number": df["recall_number"],
        "Event ID": df["event_id"],
        "Report Date": df["report_date"].dt.date,
        "Year": df["year"],
        "Firm": df["recalling_firm"],
        "Product": df["product_description"].fillna("").str.slice(0, 120),
        "Route": df["route"].fillna(""),
        "Injectable": df["injectable"],
        "Class": df["classification"],
        "Status": df["status"],
        "Reason for Recall": df["reason_for_recall"],
        "Defect Category": df["defect_category"],
        "Category Rationale": df["category_rationale"],
        "Reviewed Category": "",
    })[LOG_COLUMNS]


def existing_reviews():
    """Reviewed Category values already in the workbook, keyed by recall number."""
    if not Path(WORKBOOK).exists():
        return {}
    ws = load_workbook(WORKBOOK)["Recall Log"]
    headers = [c.value for c in ws[1]]
    num, rev = headers.index("Recall Number"), headers.index("Reviewed Category")
    return {row[num]: row[rev] for row in ws.iter_rows(min_row=2, values_only=True) if row[rev]}


def style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    log = load_log()
    reviews = existing_reviews()
    log["Reviewed Category"] = log["Recall Number"].map(reviews).fillna("")
    if reviews:
        print(f"kept {len(reviews)} reviewed categories")

    wb = Workbook()
    ws = wb.active
    ws.title = "Recall Log"
    ws.append(LOG_COLUMNS)
    for row in log.itertuples(index=False):
        ws.append(list(row))
    style_header(ws)

    for i, name in enumerate(LOG_COLUMNS, start=1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = WIDTHS[name]
        for cell in ws[col][1:]:
            cell.alignment = Alignment(wrap_text=name in WRAPPED, vertical="top")
            if name == "Report Date":
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    cs = wb.create_sheet("Categories")
    cs.append(["Category", "Definition"])
    for name, definition in CATEGORIES:
        cs.append([name, definition])
    style_header(cs)
    cs.column_dimensions["A"].width = 36
    cs.column_dimensions["B"].width = 100
    for row in cs.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    cs.freeze_panes = "A2"

    add_dropdowns(wb)
    wb.save(WORKBOOK)
    print(f"saved {len(log)} recalls to {WORKBOOK}")


if __name__ == "__main__":
    main()
